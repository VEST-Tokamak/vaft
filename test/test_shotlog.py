"""The ShotLog port: parsing, the modern card, per-shot records, FileDB (#995)."""

from __future__ import annotations

import json
from pathlib import Path
import unicodedata

import pytest

from _shotlog_fixtures import legacy_workbook, modern_workbook
from vaft.database.filedb import FileDB, FileDBPathError
from vaft.machine_mapping.pulse_schedule import (
    archive_workbooks,
    build_shot_records,
    convert_directory,
    convert_sheet,
    discover_sources,
    packaged_registry,
    record_path,
    trigger_table,
    write_extraction,
)
from vaft.machine_mapping.pulse_schedule.values import (
    normalize_status,
    parse_position,
    parse_time_window,
    parse_valve,
    parse_value,
)

REGISTRY = packaged_registry()


# --------------------------------------------------------------------------- #
# values
# --------------------------------------------------------------------------- #


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("490-510", (490, 510)),
        ("300-301.2", (300, 301.2)),
        ("400-401 ms", (400, 401)),
        ("900~902ms", (900, 902)),
        ("490", (490, 490)),
        (285, (285, 285)),
        ("510-490", None),  # reversed is a typo, not a window
        ("300-303.5\n490-", None),
        ("TS", None),
        (None, None),
    ],
)
def test_time_windows_are_parsed_or_refused(raw, expected):
    assert parse_time_window(raw) == expected


def test_value_records_keep_raw_text_and_derive_what_they_can():
    assert parse_value("12*3")["derived_value"] == 36
    assert parse_value("15*10/2")["derived_value"] == 75
    ranged = parse_value("900~902", trigger=True)
    assert (ranged["trigger_onset"], ranged["trigger_offset"]) == (900, 902)
    assert parse_value("A (OH)")["validation_flags"] == ["not_numeric"]


def test_valve_and_position_settings():
    assert parse_valve("90V(H2)") == {"voltage_V": 90, "species": "H2"}
    assert parse_valve("90 V (He)") == {"voltage_V": 90, "species": "He"}
    assert parse_valve("90(H2)") == {"voltage_V": 90, "species": "H2"}
    assert parse_valve("Low Field Side") is None
    assert parse_position("0.7/0.55") == [0.7, 0.55]


def test_status_never_claims_success_nobody_wrote_down():
    assert normalize_status(["Remarks", "fail: no plasma"])["normalized"] == "failure"
    assert normalize_status(["Remarks", "Fail"])["normalized"] == "unknown"  # headers are not evidence
    assert normalize_status(["data not saved"])["normalized"] == "not_recorded"


# --------------------------------------------------------------------------- #
# the modern card
# --------------------------------------------------------------------------- #


@pytest.fixture
def modern(tmp_path):
    return modern_workbook(
        tmp_path / "ShotLog_2025_09 #46590-46599.xlsx",
        {"250915": [
            {"shot": 46590, "tf": 160, "diagnostics": {"IF": "490-510", "SXR1": 485, "TP": 0.7},
             "ech": "100-450", "nbi_t0": 494, "nbi_dt": 3, "hi_spark": 480,
             "gas": ("90V(H2)", "300-302"), "ip": 83, "ip_length": 24, "remark": "good"},
            # Blank diagnostic cells: the IF/SXR1 settings carry forward.
            {"shot": 46591, "gas": ("95V(H2)", "300-303"), "ip": 51},
            {"shot": 46592, "diagnostics": {"IF": "480-500"}, "remark": "fail: no breakdown"},
        ]},
        title="Coil test",
    )


def _shots(dataset):
    return {shot["shot"]: shot for group in dataset["run_groups"] for shot in group["shots"]}


def test_a_modern_card_is_not_mistaken_for_the_older_integrated_template(modern):
    # The card also carries every integrated_v3 header (C (mF), NBI, gas
    # injection, remark); a tie that went the wrong way lost every trigger.
    assert convert_sheet(modern, "250915", REGISTRY)["schema_version"] == "modern_v4"


def test_the_card_reads_values_beside_labels_not_the_next_label(modern):
    shots = _shots(convert_sheet(modern, "250915", REGISTRY))
    first = shots[46590]
    timing = {(item["system"], item["identifier"]): item for item in first["planned_configuration"]["timing"]}
    assert timing[("diagnostic", "IF")]["window_ms"] == [490, 510]
    assert timing[("diagnostic", "IF")]["cell"] == "K5"
    assert timing[("diagnostic", "SXR1")]["window_ms"] == [485, 485]
    assert timing[("ec", "ECH (2.45 GHz)")]["window_ms"] == [100, 450]
    assert timing[("nbi", "NBI T0 (ms)")]["window_ms"] == [494, 494]
    assert timing[("hi", "HI PFNspark")]["window_ms"] == [480, 480]
    gas = timing[("gas", "LFS")]
    assert (gas["window_ms"], gas["valve_voltage_V"], gas["species"]) == ([300, 302], 90, "H2")
    assert first["planned_configuration"]["probe_positions"]["TP"]["positions_m"] == [0.7]
    assert first["planned_configuration"]["heating_and_current_drive"]["nbi_duration_ms"] == 3
    # No header or label text ever lands in a value slot (the #995 defect).
    raws = {item["raw"] for item in first["planned_configuration"]["timing"]}
    assert not raws & {"TS", "RF (9 dBm)", "ECH (8 GHz)", "Remarks", "Low Field Side"}


def test_plasma_current_belongs_to_the_card_whose_header_row_holds_it(modern):
    """Ip is logged on the card's own second header row, above its shot number.

    Read by block (shot row to next shot row) it would land on the *previous*
    shot, and the header-alias search read "Remarks" as its value.
    """
    shots = _shots(convert_sheet(modern, "250915", REGISTRY))
    assert shots[46590]["observed_outcome"]["plasma_current"]["peak"]["derived_value"] == 83
    assert shots[46590]["observed_outcome"]["plasma_current"]["pulse_length"]["derived_value"] == 24
    assert shots[46591]["observed_outcome"]["plasma_current"]["peak"]["derived_value"] == 51
    assert "peak" not in shots[46592]["observed_outcome"]["plasma_current"]
    assert shots[46590]["observed_outcome"]["remarks"] == ["good"]
    assert shots[46592]["status"]["normalized"] == "failure"


def test_a_card_is_recognised_by_its_section_titles_not_column_a(tmp_path):
    # 2023-06-24: the operator overtyped "Shot#" with "C"; 2025-01-23 left it empty.
    source = modern_workbook(tmp_path / "ShotLog_2023_06 #39860-39861.xlsx", {"230624": [
        {"shot": 39860, "header_a": "C", "diagnostics": {"SXR1": "500-501", "TP": 0.8}},
        {"shot": 39861, "header_a": None, "diagnostics": {"TP": 0.57}},
    ]})
    shots = _shots(convert_sheet(source, "230624", REGISTRY))
    timing = shots[39860]["planned_configuration"]["timing"]
    assert [(item["identifier"], item["window_ms"]) for item in timing] == [("SXR1", [500, 501])]
    assert shots[39861]["planned_configuration"]["probe_positions"]["TP"]["positions_m"] == [0.57]


def test_a_reference_card_is_a_reference_not_the_previous_shots_settings(tmp_path):
    """2024-08-21: a card headed "REF!!! 43013" follows shot 43483.

    Its triggers are 43013's. Read as part of 43483's block, they were once
    filed as 43483's own and carried forward through the rest of the day.
    """
    source = modern_workbook(tmp_path / "ShotLog_2024_08 #43483-43484.xlsx", {"240821": [
        {"shot": 43483, "gas": ("150V", "300-305")},
        {"shot": "REF!!!\n43013", "diagnostics": {"TS": 326, "PIG": 500}},
        {"shot": 43484},
    ]})
    dataset = convert_sheet(source, "240821", REGISTRY)
    shots = _shots(dataset)
    assert [item["identifier"] for item in shots[43483]["planned_configuration"]["timing"]] == ["LFS"]
    assert shots[43484]["planned_configuration"].get("timing", []) == []
    references = [ref["shot"] for group in dataset["run_groups"] for ref in group["references"]]
    assert references == [43013]


def test_one_card_can_log_several_shots_and_a_bad_one_is_reported(tmp_path):
    source = modern_workbook(tmp_path / "ShotLog_2023_06 #39853-39856.xlsx", {"230624": [
        {"shot": "39853\n39854\n39855", "diagnostics": {"IF": "490-510"}},
        {"shot": "41538-41440"},  # reversed: unreadable, not silently skipped
        {"shot": "41560\n~41561"},  # a range broken across lines is still one range
        {"shot": 39856},
    ]})
    dataset = convert_sheet(source, "230624", REGISTRY)
    shots = _shots(dataset)
    assert sorted(shots) == [39853, 39854, 39855, 39856, 41560, 41561]
    assert shots[39855]["planned_configuration"]["timing"][0]["window_ms"] == [490, 510]
    assert [item["raw"] for item in dataset["review"]["unreadable_shot_cells"]] == ["41538-41440"]


def test_a_title_between_cards_opens_a_run_group_and_card_rows_do_not(tmp_path):
    source = modern_workbook(tmp_path / "ShotLog_2024_08 #43481-43484.xlsx", {"240821": [
        {"shot": 43481},
        {"shot": 43482},
        # A reference card's two-cell rows ("VB", "filter") are not titles.
        {"shot": "REF!!!\n43013", "diagnostics": {"TS": 326}},
        {"shot": 43483, "title_before": "오믹 방전 만들기/ BZn"},
        {"shot": 43484},
    ]}, title="Coil 전류 테스트")
    dataset = convert_sheet(source, "240821", REGISTRY)
    groups = [(group["title"]["raw"], [shot["shot"] for shot in group["shots"]]) for group in dataset["run_groups"]]
    assert groups == [("Coil 전류 테스트", [43481, 43482]), ("오믹 방전 만들기/ BZn", [43483, 43484])]


def test_session_documents_are_reproducible(modern):
    assert convert_sheet(modern, "250915", REGISTRY) == convert_sheet(modern, "250915", REGISTRY)


def test_legacy_sheets_still_convert(tmp_path):
    source = legacy_workbook(tmp_path / "ShotLog_2014_01.xlsx")
    dataset = convert_sheet(source, "20140117", REGISTRY)
    assert dataset["schema_version"] == "legacy_v1"
    assert dataset["experiment_session"]["experiment_date"] == "2014-01-17"
    shot = _shots(dataset)[7344]
    assert shot["planned_configuration"]["magnetic"]["tf"]["derived_value"] == 10
    assert shot["status"]["normalized"] == "failure"


def test_overrides_are_yaml_and_bound_to_the_workbook_bytes(tmp_path, modern):
    dataset = convert_sheet(modern, "250915", REGISTRY)
    group = dataset["run_groups"][0]["id"]
    overrides = tmp_path / "overrides.yaml"
    overrides.write_text(
        "- id: fix-1\n"
        f"  source_sha256: {dataset['provenance']['source_sha256']}\n"
        f"  target: {{session_id: '{dataset['experiment_session']['id']}', run_group_id: '{group}', shot: 46591}}\n"
        "  path: status.normalized\n"
        "  value: failure\n"
        "  reason: operator note\n"
        "- id: other-workbook\n"
        "  source_sha256: deadbeef\n"
        "  target: {session_id: x, run_group_id: y, shot: 1}\n"
        "  path: status.normalized\n"
        "  value: failure\n",
        encoding="utf-8",
    )
    patched = convert_sheet(modern, "250915", REGISTRY, overrides)
    assert _shots(patched)[46591]["status"]["normalized"] == "failure"
    assert patched["review"]["applied_overrides"] == ["fix-1"]
    # An override for another workbook is not a conflict of this one.
    assert patched["review"]["override_conflicts"] == []


# --------------------------------------------------------------------------- #
# discovery
# --------------------------------------------------------------------------- #


def test_every_file_gets_a_role_and_only_residue_is_left_out(tmp_path):
    # macOS lists Korean names decomposed (NFD); an autosave spelled that way
    # once passed every exclusion and was archived as its month's record.
    autosave = unicodedata.normalize("NFD", "ShotLog_2016_09 #15742-15822 (자동 저장됨).xlsx")
    (tmp_path / autosave).write_bytes(b"")
    (tmp_path / "ShotLog_2016_09 #15742-15822.xlsx").write_bytes(b"")
    for name in (
        "ShotLog_2026_02 #47752-.xlsx",          # only record of February: converted
        "ShotLog_2026_03 #47986-.xlsx",          # superseded by the closed one: kept, not converted
        "ShotLog_2026_03 #47986-48192.xlsx",
        "~$ShotLog_2026_03 #47986-48192.xlsx",
        "._ShotLog_2026_03 #47986-48192.xlsx",
        "복사본 ShotLog_2024_5 #42213-.xlsx",
        "blank.xlsx",                            # no shot numbers: kept, not converted
    ):
        (tmp_path / name).write_bytes(b"")
    (tmp_path / "Shot Log form").mkdir()
    (tmp_path / "Shot Log form" / "VEST shotlog_New_v3.xlsx").write_bytes(b"")
    legacy_workbook(tmp_path / "ERC_ShotLog.xlsx", shots=(2774, 2776))   # a real log outside the monthly naming
    (tmp_path / "ShotLog_dataset").mkdir()           # the standalone tool's output: never a source
    (tmp_path / "ShotLog_dataset" / "batch-manifest.yaml").write_text("x")

    discovery = discover_sources(tmp_path)
    assert [(item.path.name, item.role) for item in discovery.included] == [
        ("ShotLog_2016_09 #15742-15822.xlsx", "monthly_record"),
        ("ShotLog_2026_02 #47752-.xlsx", "monthly_record"),
        ("ShotLog_2026_03 #47986-48192.xlsx", "monthly_record"),
        ("ERC_ShotLog.xlsx", "supplementary_log"),
    ]
    # The open-ended February is bounded by March's first shot.
    assert discovery.included[1].span() == (47752 - 200, 47986 + 200)
    kept = {item.archive_path: item.role for item in discovery.archived_only}
    assert kept == {
        "other/ShotLog_2026_03 #47986-.xlsx": "superseded_open_ended",
        "other/" + unicodedata.normalize("NFC", autosave): "autosave",
        "other/복사본 ShotLog_2024_5 #42213-.xlsx": "copy",
        "other/Shot Log form/VEST shotlog_New_v3.xlsx": "template",
        "other/blank.xlsx": "unrecognised_workbook",
    }
    reasons = {item["source_file"]: item["reason"] for item in discovery.excluded}
    assert reasons == {
        "~$ShotLog_2026_03 #47986-48192.xlsx": "excel_lock_file",
        "._ShotLog_2026_03 #47986-48192.xlsx": "appledouble_sidecar",
        # Skipped, but never silently: the listing says what and why.
        "ShotLog_dataset (derived output of the standalone tool)": "skipped_directory",
    }


def test_role_words_match_whole_words_only(tmp_path):
    # "formation" contains "form", "Threshold" contains "old": neither hides a month.
    (tmp_path / "Threshold scans").mkdir()
    for relative in ("ShotLog_2024_06 #47000-47300 (plasma formation).xlsx",
                     "Threshold scans/ShotLog_2024_07 #47301-47400.xlsx",
                     "ShotLog_2016_10 #15823-old.xlsx", "ShotLog_2015_02 #11314- (temp).xlsx"):
        (tmp_path / relative).write_bytes(b"")
    discovery = discover_sources(tmp_path)
    assert [item.path.name for item in discovery.included] == [
        "ShotLog_2024_06 #47000-47300 (plasma formation).xlsx", "ShotLog_2024_07 #47301-47400.xlsx",
    ]
    assert {item.role for item in discovery.archived_only} == {"copy", "autosave"}


def test_two_files_with_one_name_keep_distinct_archive_paths(tmp_path):
    for folder, shot in (("a", 40000), ("b", 40001)):
        (tmp_path / folder).mkdir()
        legacy_workbook(tmp_path / folder / "ShotLog_2023_05 #40000-40300.xlsx", sheet="230501", shots=(shot,))
    paths = [item.archive_path for item in discover_sources(tmp_path).included]
    assert len(set(paths)) == 2 and paths[0] == "2023/ShotLog_2023_05 #40000-40300.xlsx"
    filedb = FileDB(tmp_path / "filedb")
    archive_workbooks(tmp_path, filedb)
    assert archive_workbooks(tmp_path, filedb)["replaced"] == []   # stable, not swapping


def test_a_typo_span_is_read_as_open_ended_not_as_an_empty_month(tmp_path):
    (tmp_path / "ShotLog_2025_03 #47000-4730.xlsx").write_bytes(b"")
    (item,) = discover_sources(tmp_path).included
    assert (item.first_shot, item.last_shot) == (47000, None)


def test_the_archive_reads_back_by_location_and_is_never_its_own_source(tmp_path, modern):
    (tmp_path / "복사본 ShotLog_2025_09 #46590-.xlsx").write_bytes(modern.read_bytes())
    legacy_workbook(tmp_path / "ERC_ShotLog.xlsx", sheet="2013-01-21", shots=(2774, 2776))
    filedb = FileDB(tmp_path / "filedb")
    archive_workbooks(tmp_path, filedb)
    # A second archive of the same folder does not ingest legacy/shotlog itself.
    assert archive_workbooks(tmp_path, filedb)["copied"] == []
    back = discover_sources(tmp_path / "filedb/legacy/shotlog/input")
    assert [(item.archive_path, item.role) for item in back.included] == [
        ("2025/ShotLog_2025_09 #46590-46599.xlsx", "monthly_record"),
        ("supplementary/ERC_ShotLog.xlsx", "supplementary_log"),
    ]
    assert [item.archive_path for item in back.fallback] == ["other/복사본 ShotLog_2025_09 #46590-.xlsx"]


def test_a_copy_speaks_only_for_shots_nothing_else_records(tmp_path):
    modern_workbook(tmp_path / "ShotLog_2015_10 #13900-13916.xlsx", {"151030": [{"shot": 13910}]})
    modern_workbook(tmp_path / "ShotLog_2015_10 #13900- (자동 저장됨).xlsx",
                    {"151030": [{"shot": 13910, "diagnostics": {"IF": "400-401"}}, {"shot": 13911}]})
    _, records, _ = _records(tmp_path)
    assert records[13910]["source"]["workbook"] == "ShotLog_2015_10 #13900-13916.xlsx"
    assert records[13910]["from_fallback"] is False and len(records[13910]["occurrences"]) == 1
    assert records[13911]["from_fallback"] is True
    assert records[13911]["occurrences"][0]["role"] == "autosave"


def test_a_log_without_a_named_span_ignores_strays_far_from_its_shots(tmp_path):
    legacy_workbook(tmp_path / "conditioning.xlsx", sheet="Sheet1",
                    shots=(4861, 4862, 4863, 4864, 4865, 4866, 4867, 4868, 4869, 400))
    (item,) = discover_sources(tmp_path).included
    lower, upper = item.span()
    assert lower <= 4861 and upper >= 4869 and not lower <= 400


def test_a_supplementary_log_is_converted_but_ranks_below_the_month(tmp_path):
    legacy_workbook(tmp_path / "ERC_ShotLog.xlsx", shots=(7344, 7345))
    source = legacy_workbook(tmp_path / "ShotLog_2014_01 #7300-7400.xlsx")
    _, records, _ = _records(tmp_path)
    record = records[7344]
    assert record["source"]["workbook"] == source.name
    assert [item["role"] for item in record["occurrences"]] == ["monthly_record", "supplementary_log"]


def test_an_unclassified_sheet_still_gives_its_shots_records(tmp_path):
    from openpyxl import Workbook

    workbook = Workbook()
    ws = workbook.active
    ws.title = "210512"
    for row, values in enumerate([[31000, "Coil", "setting"], [None, "x", 1], [31001, "Coil", "setting"]], start=1):
        for column, value in enumerate(values, start=1):
            if value is not None:
                ws.cell(row, column, value)
    path = tmp_path / "ShotLog_2021_05 #31000-31001.xlsx"
    workbook.save(path)
    dataset = convert_sheet(path, "210512", REGISTRY)
    assert dataset["schema_version"] == "unclassified"
    assert sorted(_shots(dataset)) == [31000, 31001]
    assert any(cell["raw"] == "Coil" for cell in _shots(dataset)[31000]["source_occurrences"][0]["cells"])
    _, records, _ = _records(tmp_path)
    assert records[31001]["schema_version"] == "unclassified"


def test_the_split_switch_power_supply_sheets_have_a_schema(tmp_path):
    from openpyxl import Workbook

    workbook = Workbook()
    ws = workbook.active
    ws.title = "190512"
    for column, value in enumerate([22462, "TF", "P/S", "PF", "SW(+)", "SW(-)", "Charging Voltage"], start=1):
        ws.cell(1, column, value)
    ws.cell(2, 2, 144)
    path = tmp_path / "ShotLog_2019_05 #22271-22500.xlsx"
    workbook.save(path)
    dataset = convert_sheet(path, "190512", REGISTRY)
    assert dataset["schema_version"] == "ps_v3"
    assert _shots(dataset)[22462]["planned_configuration"]["magnetic"]["tf"]["derived_value"] == 144


def test_a_number_outside_the_workbook_span_is_not_a_shot(tmp_path):
    source = modern_workbook(tmp_path / "ShotLog_2025_09 #46590-46599.xlsx", {"250915": [
        {"shot": 43960, "diagnostics": {"SXR1": "485-535"}},   # an older shot quoted as a reference
        {"shot": 46590},
        {"shot": 1000},                                        # not a shot at all
    ]})
    span = (46590 - 200, 46599 + 200)
    dataset = convert_sheet(source, "250915", REGISTRY, span=span)
    assert sorted(_shots(dataset)) == [46590]
    assert [item["shots"] for item in dataset["review"]["out_of_span_markers"]] == [[43960], [1000]]
    references = [ref["shot"] for group in dataset["run_groups"] for ref in group["references"]]
    assert references == [43960, 1000]
    # The quoted card's settings did not become 46590's.
    assert _shots(dataset)[46590]["planned_configuration"].get("timing", []) == []


def test_coverage_names_the_gaps_and_where_to_look(tmp_path):
    from vaft.machine_mapping.pulse_schedule.records import coverage_report

    modern_workbook(tmp_path / "ShotLog_2025_09 #46590-46599.xlsx", {"250915": [
        {"shot": 46590}, {"shot": 46591}, {"shot": 46595},
    ]})
    _, records, _ = _records(tmp_path)
    report = coverage_report(records)
    assert (report["first_logged_shot"], report["no_source_before"]) == (46590, 46590)
    assert report["missing_ranges"] == [{
        "first": 46592, "last": 46594, "count": 3,
        "before": {"shot": 46591, "workbook": "ShotLog_2025_09 #46590-46599.xlsx", "sheet": "250915"},
        "after": {"shot": 46595, "workbook": "ShotLog_2025_09 #46590-46599.xlsx", "sheet": "250915"},
    }]


# --------------------------------------------------------------------------- #
# per-shot records
# --------------------------------------------------------------------------- #


def _records(directory):
    sessions, manifest = convert_directory(directory)
    return sessions, build_shot_records(sessions), manifest


def test_diagnostic_triggers_carry_forward_and_say_so(tmp_path, modern):
    _, records, _ = _records(tmp_path)
    second = {(item["system"], item["identifier"]): item for item in records[46591]["effective_timing"]}
    assert second[("diagnostic", "IF")]["window_ms"] == [490, 510]
    assert second[("diagnostic", "IF")]["source_shot"] == 46590
    assert second[("diagnostic", "IF")]["inherited"] is True
    # Heating and gas are copied onto every card, so they are never inherited.
    assert ("ec", "ECH (2.45 GHz)") not in second
    assert second[("gas", "LFS")]["window_ms"] == [300, 303]
    third = {(item["system"], item["identifier"]): item for item in records[46592]["effective_timing"]}
    assert third[("diagnostic", "IF")]["window_ms"] == [480, 500]
    assert third[("diagnostic", "IF")]["inherited"] is False
    assert records[46592]["effective_probe_positions"]["TP"]["source_shot"] == 46590


def test_the_trigger_table_is_on_the_daq_clock_in_the_shape_sxr_reads(tmp_path, modern):
    _, records, _ = _records(tmp_path)
    table = trigger_table(records)["shots"]
    assert table[46590]["IF"] == {"start_time_ms": 290, "end_time_ms": 310, "source_shot": 46590}
    # Filed as SXR, the key soft_x_rays reads for the primary array.
    assert table[46591]["SXR"] == {"start_time_ms": 285, "end_time_ms": 285, "source_shot": 46590}
    assert "SXR1" not in table[46591]
    assert table[46590]["TP"] == {"measured_position_m": [0.7], "source_shot": 46590}


def test_a_shots_own_sxr1_beats_an_inherited_bare_sxr_in_the_table(tmp_path):
    modern_workbook(tmp_path / "ShotLog_2025_09 #46590-46591.xlsx", {"250915": [
        {"shot": 46590, "diagnostics": {"SXR1": "500-501"}},
        {"shot": 46591, "diagnostics": {"SXR1": "485-535"}},
    ]})
    _, records, _ = _records(tmp_path)
    # Simulate the other spelling on the earlier card: a bare SXR carried in.
    carried = dict(records[46591]["effective_timing"][0], identifier="SXR", inherited=True, source_shot=46590,
                   window_ms=[500, 501])
    records[46591]["effective_timing"].append(carried)
    assert trigger_table(records)["shots"][46591]["SXR"]["start_time_ms"] == 285


def test_a_shot_in_two_workbooks_takes_the_one_whose_span_covers_it(tmp_path):
    modern_workbook(tmp_path / "ShotLog_2025_08 #46500-46589.xlsx",
                    {"250831": [{"shot": 46590, "diagnostics": {"IF": "470-490"}}]})
    modern_workbook(tmp_path / "ShotLog_2025_09 #46590-46599.xlsx",
                    {"250901": [{"shot": 46590, "diagnostics": {"IF": "490-510"}}]})
    _, records, _ = _records(tmp_path)
    record = records[46590]
    assert record["source"]["workbook"] == "ShotLog_2025_09 #46590-46599.xlsx"
    assert len(record["occurrences"]) == 2
    assert record["ambiguous"] is True


# --------------------------------------------------------------------------- #
# FileDB
# --------------------------------------------------------------------------- #


def test_the_grammar_admits_a_diagnostic_scoped_legacy_artifact():
    db = FileDB("/filedb")
    assert db.legacy("shotlog", None, artifact="input") == Path("/filedb/legacy/shotlog/input")
    assert db.legacy("shotlog", 46590, artifact="metadata") == Path("/filedb/legacy/shotlog/46590/metadata")
    with pytest.raises(FileDBPathError, match="must name an artifact"):
        db.legacy("shotlog", None)


def test_archive_copies_verifies_and_keeps_superseded_bytes(tmp_path, modern):
    filedb = FileDB(tmp_path / "filedb")
    first = archive_workbooks(tmp_path, filedb)
    stored = tmp_path / "filedb/legacy/shotlog/input/2025" / modern.name
    assert first["copied"] == [f"2025/{modern.name}"] and stored.read_bytes() == modern.read_bytes()
    assert archive_workbooks(tmp_path, filedb)["unchanged"] == [f"2025/{modern.name}"]

    modern_workbook(modern, {"250915": [{"shot": 46590, "diagnostics": {"IF": "480-500"}}]})
    third = archive_workbooks(tmp_path, filedb)
    assert third["replaced"] == [f"2025/{modern.name}"]
    manifest = json.loads((tmp_path / "filedb/legacy/shotlog/input/manifest.json").read_text())
    kept = tmp_path / "filedb/legacy/shotlog/input" / manifest["history"][0]["kept_as"]
    assert kept.exists() and manifest["history"][0]["kept_as"].startswith("superseded/2025/")
    assert manifest["files"][f"2025/{modern.name}"]["role"] == "monthly_record"
    # The archive reads back as a source folder, and superseded bytes are not rediscovered.
    assert [item.path for item in discover_sources(tmp_path / "filedb/legacy/shotlog/input").included] == [stored]


def test_archive_keeps_copies_and_extra_files_byte_for_byte(tmp_path, modern):
    (tmp_path / "복사본 ShotLog_2025_09 #46590-.xlsx").write_bytes(modern.read_bytes())
    elsewhere = tmp_path.parent / f"{tmp_path.name}-outside"
    elsewhere.mkdir()
    stray = legacy_workbook(elsewhere / "ShotLog_2022_02 #35060-.xlsx")
    filedb = FileDB(tmp_path / "filedb")
    summary = archive_workbooks(tmp_path, filedb, extra=[stray])
    root = tmp_path / "filedb/legacy/shotlog/input"
    assert (root / "other/복사본 ShotLog_2025_09 #46590-.xlsx").read_bytes() == modern.read_bytes()
    assert (root / "2022" / stray.name).read_bytes() == stray.read_bytes()
    assert len(summary["copied"]) == 3
    # The copy is kept, never converted.
    assert [item.role for item in discover_sources(root).included] == ["monthly_record", "monthly_record"]


def test_extraction_writes_records_and_touches_nothing_on_a_rerun(tmp_path, modern):
    filedb = FileDB(tmp_path / "filedb")
    sessions, records, manifest = _records(tmp_path)
    counts = write_extraction(sessions, records, manifest, filedb)
    assert counts["records_written"] == 3
    record = json.loads(record_path(filedb, 46590).read_text())
    assert record["shot"] == 46590 and record["clock"]["daq_offset_ms"] == -200
    assert (tmp_path / "filedb/legacy/shotlog/output/experiment-days/2025/2025-09-15__250915.yaml").exists()
    coverage = json.loads((tmp_path / "filedb/legacy/shotlog/output/coverage.json").read_text())
    assert (coverage["records"], coverage["missing_count"]) == (3, 0)
    sessions, records, manifest = _records(tmp_path)
    assert write_extraction(sessions, records, manifest, filedb)["records_unchanged"] == 3


def test_the_cli_extracts_then_builds_the_trigger_table_from_the_stored_records(tmp_path, modern, capsys):
    import yaml

    from vaft.cli._main import main as cli_main

    root = tmp_path / "filedb"
    assert cli_main(["shotlog", "archive", "--source", str(tmp_path), "--filedb", str(root)]) == 0
    assert cli_main(["shotlog", "extract", "--filedb", str(root)]) == 0
    output = tmp_path / "triggers.yaml"
    assert cli_main(["shotlog", "triggers", "--filedb", str(root), "--output", str(output)]) == 0
    table = yaml.safe_load(output.read_text(encoding="utf-8"))["shots"]
    assert table[46591]["IF"] == {"start_time_ms": 290, "end_time_ms": 310, "source_shot": 46590}
