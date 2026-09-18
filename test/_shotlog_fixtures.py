"""Synthetic ShotLog workbooks laid out like the real templates (#995).

The real workbooks live on an operator's disk and in FileDB, never in the
repository, so the tests build the layouts they exercise. The modern card
below reproduces the 2023 template cell for cell (columns J..W), including
the two traps it sets: the plasma current sits on the card's second header
row, and every value slot under a header holds the *next label*.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook

CARD_HEIGHT = 8  # two header rows + six card rows


def write_card(ws: Any, top: int, shot: int, **values: Any) -> int:
    """Write one card with its header at row ``top``; return the next free row."""
    h1, h2, s = top, top + 1, top + 2
    cell = ws.cell
    # ``shot`` may be a label such as "REF!!!\n43013"; ``header_a`` is what
    # the operator left in the header's column A (normally "Shot#").
    if values.get("header_a", "Shot#") is not None:
        cell(h1, 1, values.get("header_a", "Shot#"))
    for column, text in ( (2, "Magnetic coil & Power system"), (10, "Diagnostics trigger (ms, m)"),
                         (14, "H & CD trigger (ms)"), (20, "Gas injection (ms)"), (22, "Ip, maximum (kA)"),
                         (23, "Ip, pulse length (ms)")):
        cell(h1, column, text)
    for column, text in ((2, "TF"), (3, "PF Coil"), (4, "P/S"), (5, "SW (ms)"), (7, "Charging V (kV)"),
                         (9, "C (mF)"), (14, "ECH (2.45 GHz)"), (15, "NBI"), (17, "HI"), (20, "Valve (V)"),
                         (21, "Trigger (ms)")):
        cell(h2, column, text)
    cell(s, 1, shot)
    if values.get("tf") is not None:
        cell(s, 2, values["tf"])
    # Diagnostic labels; the value is the cell to the right.
    for offset, label in enumerate(("TS", "IF", "IDS", "HXR", "SXR1", "SXR2")):
        cell(s + offset, 10, label)
    for offset, label in enumerate(("PIG", "TP", "CES", "MD", "FastCam")):
        cell(s + offset, 12, label)
    positions = {"TS": (s, 11), "IF": (s + 1, 11), "IDS": (s + 2, 11), "HXR": (s + 3, 11),
                 "SXR1": (s + 4, 11), "SXR2": (s + 5, 11), "PIG": (s, 13), "TP": (s + 1, 13),
                 "CES": (s + 2, 13), "FastCam": (s + 4, 13)}
    for label, value in values.get("diagnostics", {}).items():
        cell(*positions[label], value)
    # H & CD: ECH window under its header, then more labels in the same column.
    if values.get("ech") is not None:
        cell(s, 14, values["ech"])
    cell(s + 1, 14, "ECH (8 GHz)")
    cell(s + 2, 14, "HV (9 kV)")
    cell(s + 4, 14, "RF (9 dBm)")
    cell(s, 15, "T0 (ms)")
    cell(s + 1, 15, "Gas_line")
    cell(s + 2, 15, "dT (ms)")
    if values.get("nbi_t0") is not None:
        cell(s, 16, values["nbi_t0"])
    if values.get("nbi_dt") is not None:
        cell(s + 2, 16, values["nbi_dt"])
    cell(s, 17, "VPFN")
    cell(s, 18, "V")
    cell(s + 1, 17, "PFNspark")
    cell(s + 1, 18, "ms")
    if values.get("hi_spark") is not None:
        cell(s + 1, 19, values["hi_spark"])
    # Gas: side label, then the valve setting with its trigger to the right.
    cell(s, 20, "Low Field Side")
    if values.get("gas") is not None:
        valve, window = values["gas"]
        cell(s + 1, 20, valve)
        cell(s + 1, 21, window)
    cell(s + 3, 20, "High Field side")
    # Plasma current on the card's own second header row.
    if values.get("ip") is not None:
        cell(h2, 22, values["ip"])
    if values.get("ip_length") is not None:
        cell(h2, 23, values["ip_length"])
    cell(s, 22, "Remarks")
    if values.get("remark") is not None:
        cell(s + 1, 22, values["remark"])
    return top + CARD_HEIGHT


def modern_workbook(path: Path, sheets: dict[str, list[dict[str, Any]]], title: str | None = None) -> Path:
    """``sheets``: ``{sheet_name: [card kwargs incl. shot]}``."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, cards in sheets.items():
        ws = workbook.create_sheet(name)
        row = 1
        if title:
            ws.cell(1, 1, title)
            row = 2
        for card in cards:
            card = dict(card)
            title_before = card.pop("title_before", None)
            if title_before:
                ws.cell(row, 1, title_before)  # an operator's run-group title
                row += 1
            row = write_card(ws, row, **card)
    workbook.save(path)
    return path


def legacy_workbook(path: Path, sheet: str = "20140117") -> Path:
    """A 2013-era sheet: headers on row 2, one shot's values on row 3."""
    workbook = Workbook()
    ws = workbook.active
    ws.title = sheet
    headers = ["Shot", "TF", "PF", "SW1", "C1", "gas", "Fail"]
    for column, text in enumerate(headers, start=1):
        ws.cell(2, column, text)
    for column, value in enumerate([7344, 10, 3, 1, 12, "900~902ms", "fail: no plasma"], start=1):
        ws.cell(3, column, value)
    workbook.save(path)
    return path
