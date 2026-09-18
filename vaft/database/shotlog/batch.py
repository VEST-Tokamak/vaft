"""Discover the monthly ShotLog workbooks and convert them all.

The ShotLog folder accumulates Excel lock files, copies (``복사본``/``사본``),
forms and in-progress workbooks next to the real ones. Discovery decides which
workbook is *the* record for each month and says why every other file was left
out, so the exclusion list is reviewable rather than silent.
"""

from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path
import re
from typing import Any, Iterator
import unicodedata

from openpyxl import load_workbook

from .converter import convert_sheet, session_path, sha256_file, validate_dataset
from .schema import packaged_registry, schema_versions

MONTHLY_RE = re.compile(
    r"^ShotLog_(?P<year>(?:19|20)\d{2})_(?P<month>0?[1-9]|1[0-2])(?:\s|#|$)", re.IGNORECASE
)
SHOT_SPAN_RE = re.compile(r"#\s*(?P<first>\d+)\s*[-~]\s*(?P<last>\d+)?\s*$")
EXCLUDED_NAME_TOKENS = ("복사본", "사본", "temp", "plan", "old", "자동저장", "자동 저장", "form")


@dataclass(frozen=True)
class SourceWorkbook:
    """One workbook chosen as the record for its month."""

    path: Path
    year: int
    month: int
    first_shot: int | None
    last_shot: int | None

    @property
    def open_ended(self) -> bool:
        return self.first_shot is not None and self.last_shot is None

    def covers(self, shot: int) -> bool:
        if self.first_shot is None:
            return False
        return self.first_shot <= shot and (self.last_shot is None or shot <= self.last_shot)


@dataclass
class Discovery:
    included: list[SourceWorkbook] = field(default_factory=list)
    excluded: list[dict[str, Any]] = field(default_factory=list)


def display_name(path: Path) -> str:
    return unicodedata.normalize("NFC", path.name)


def _describe(path: Path) -> SourceWorkbook | None:
    match = MONTHLY_RE.match(path.stem)
    if not match:
        return None
    span = SHOT_SPAN_RE.search(path.stem)
    first = int(span.group("first")) if span else None
    last = int(span.group("last")) if span and span.group("last") else None
    return SourceWorkbook(path, int(match.group("year")), int(match.group("month")), first, last)


def classify_source(path: Path) -> tuple[SourceWorkbook | None, str | None]:
    """``(workbook, None)`` for a monthly ShotLog, else ``(None, reason)``.

    Names are compared in NFC. macOS hands back Korean file names decomposed
    (NFD), in which ``자동 저장`` never matches the composed token -- which is
    how two ``(자동 저장됨)`` autosaves were once taken for their months' record.
    """
    name = unicodedata.normalize("NFC", path.name)
    if name.startswith("._"):
        return None, "appledouble_sidecar"
    if name.startswith("~$"):
        return None, "excel_lock_file"
    lowered = name.casefold()
    if any(token.casefold() in lowered for token in EXCLUDED_NAME_TOKENS):
        return None, "copy_or_temporary_or_template"
    described = _describe(path)
    if described is None:
        return None, "not_monthly_shotlog"
    return described, None


def discover_sources(input_dir: str | Path) -> Discovery:
    """Choose one workbook per month from ``input_dir``.

    An open-ended workbook (``#47752-``) is the in-progress copy of a month.
    It is superseded when the same month also has a closed one (``#47752-47985``)
    and is otherwise the only record of that month, so it is kept. Dropping
    every open-ended name, as the standalone tool did, silently lost whole
    months (2026-02, 2026-04, 2026-09, ...).
    """
    discovery = Discovery()
    by_month: dict[tuple[int, int], list[SourceWorkbook]] = defaultdict(list)
    # The operators' folder is flat; the FileDB archive files by year.
    root = Path(input_dir)
    paths = [*root.glob("*.xlsx"), *root.glob("[12][0-9][0-9][0-9]/*.xlsx")]
    for path in sorted(paths, key=lambda item: item.name.casefold()):
        workbook, reason = classify_source(path)
        if workbook is None:
            discovery.excluded.append({"source_file": display_name(path), "reason": reason})
        else:
            by_month[(workbook.year, workbook.month)].append(workbook)
    for month in sorted(by_month):
        candidates = by_month[month]
        closed = [item for item in candidates if not item.open_ended]
        keep = closed or candidates
        for item in candidates:
            if item in keep:
                discovery.included.append(item)
            else:
                discovery.excluded.append({
                    "source_file": display_name(item.path),
                    "reason": "superseded_open_ended_workbook",
                    "superseded_by": [display_name(other.path) for other in closed],
                })
    return discovery


def iter_sessions(
    sources: list[SourceWorkbook],
    registry: dict[str, Any] | None = None,
    overrides_path: Path | None = None,
    manifest: dict[str, Any] | None = None,
) -> Iterator[tuple[SourceWorkbook, dict[str, Any]]]:
    """Yield ``(workbook, session)`` for every classified sheet.

    Per-workbook and per-sheet failures are recorded in ``manifest`` (when
    given) and skipped: one corrupt month must not stop the archive.
    """
    registry = registry or packaged_registry()
    record = manifest if manifest is not None else _new_manifest(registry)
    for source in sources:
        source_hash = sha256_file(source.path)
        entry: dict[str, Any] = {"source_file": display_name(source.path), "source_sha256": source_hash, "sheets": []}
        record["included_files"].append(entry)
        try:
            workbook = load_workbook(source.path, read_only=False, data_only=False)
        except Exception as error:  # noqa: BLE001 - recorded, the walk continues
            entry["error"] = repr(error)
            record["errors"].append({"source_file": display_name(source.path), "stage": "open_workbook", "error": repr(error)})
            continue
        try:
            for sheet_name in list(workbook.sheetnames):
                sheet: dict[str, Any] = {"sheet_name": sheet_name}
                entry["sheets"].append(sheet)
                try:
                    dataset = convert_sheet(source.path, sheet_name, registry, overrides_path,
                                            workbook=workbook, source_hash=source_hash)
                except Exception as error:  # noqa: BLE001
                    sheet.update({"status": "error", "error": repr(error)})
                    record["errors"].append({"source_file": display_name(source.path), "sheet_name": sheet_name,
                                             "stage": "convert_sheet", "error": repr(error)})
                    continue
                sheet["schema_version"] = dataset["schema_version"]
                if dataset["schema_version"] == "unclassified":
                    sheet["status"] = "unclassified"
                    record["unclassified"].append({"source_file": display_name(source.path), "sheet_name": sheet_name,
                                                   "classification": dataset.get("classification", {})})
                    continue
                errors = validate_dataset(dataset, registry)
                sheet.update({"status": "converted", "session_id": dataset["experiment_session"]["id"],
                              "validation_errors": errors})
                if errors:
                    record["errors"].append({"source_file": display_name(source.path), "sheet_name": sheet_name,
                                             "stage": "validation", "errors": errors})
                yield source, dataset
        finally:
            workbook.close()


def _new_manifest(registry: dict[str, Any]) -> dict[str, Any]:
    return {
        "batch_version": 2,
        "schema_versions": schema_versions(registry),
        "included_files": [],
        "excluded_files": [],
        "unclassified": [],
        "errors": [],
        "collisions": [],
    }


def summarise(manifest: dict[str, Any]) -> dict[str, int]:
    sheets = [sheet for item in manifest["included_files"] for sheet in item.get("sheets", [])]
    return {
        "xlsx_included": len(manifest["included_files"]),
        "xlsx_excluded": len(manifest["excluded_files"]),
        "sheets_converted": sum(1 for sheet in sheets if sheet.get("status") == "converted"),
        "sheets_unclassified": len(manifest["unclassified"]),
        "errors": len(manifest["errors"]),
        "collisions": len(manifest["collisions"]),
    }


def convert_directory(
    input_dir: str | Path,
    registry: dict[str, Any] | None = None,
    overrides_path: Path | None = None,
) -> tuple[list[tuple[SourceWorkbook, dict[str, Any]]], dict[str, Any]]:
    """Convert every chosen workbook; return the sessions and a batch manifest.

    Two sheets can name the same session (``250219`` in two workbooks). The
    second keeps its content under ``<session>__source-<sha8>`` and the pair is
    listed under ``collisions`` instead of one overwriting the other.
    """
    registry = registry or packaged_registry()
    discovery = discover_sources(input_dir)
    manifest = _new_manifest(registry)
    manifest["excluded_files"] = discovery.excluded
    sessions: list[tuple[SourceWorkbook, dict[str, Any]]] = []
    claimed: dict[str, tuple[str, str]] = {}
    for source, dataset in iter_sessions(discovery.included, registry, overrides_path, manifest):
        session = dataset["experiment_session"]
        original = session["id"]
        owner = claimed.get(original)
        if owner is not None:
            provenance = dataset["provenance"]
            renamed = f"{original}__source-{provenance['source_sha256'][:8]}"
            if renamed in claimed:
                renamed = f"{renamed}-{len(manifest['collisions']) + 1}"
            manifest["collisions"].append({
                "session_id": original, "renamed_to": renamed,
                "source_file": display_name(source.path), "sheet_name": provenance["sheet_name"],
                "existing_source_file": owner[0], "existing_sheet_name": owner[1],
            })
            session["id"] = renamed
            for group in dataset["run_groups"]:
                group["id"] = renamed + group["id"][len(original):]
        claimed[session["id"]] = (display_name(source.path), dataset["provenance"]["sheet_name"])
        sessions.append((source, dataset))
    manifest["summary"] = summarise(manifest)
    return sessions, manifest


__all__ = [
    "Discovery",
    "SourceWorkbook",
    "classify_source",
    "convert_directory",
    "discover_sources",
    "iter_sessions",
    "session_path",
    "summarise",
]
