"""Read one shot's card from a ``modern_v4`` ShotLog sheet.

From 2023 the ShotLog is a stack of identical cards, one per shot::

    row h1   Shot# | Magnetic coil & Power system | Diagnostics trigger | H & CD trigger | Gas injection | Ip, maximum | Ip, pulse length
    row h2         | TF | PF Coil | P/S | SW | ...  |                     | ECH | NBI | HI  | Valve (V) | Trigger (ms) | <Ip> | <length>
    row s   <shot> | ...                           | TS  <t> PIG <t>     | <ECH window> T0 (ms) <t> VPFN V <v> | Low Field Side | Remarks
    row s+1        |                               | IF  <t> TP  <m>     | ECH (8 GHz)  Gas_line ...     | 90V(H2) 300-302 | <remark>
    ...

Every value sits in a fixed position relative to a *label*, not relative to
the header: a diagnostic's trigger is the cell to the right of its name, the
2.45 GHz ECH window is the cell under its header, a gas trigger is the cell to
the right of the valve setting that follows a side label. The header-alias
search the other eras use reads the cell below a header, which on this card is
the next label -- that is how ``"TS"`` and ``"RF (9 dBm)"`` used to be stored
as values (#995).

Columns drift by one between template revisions (gas at T or S), so every
section is located from its header text on each card rather than by letter.
Plasma current is logged on row h2 under its header, i.e. on the card's own
header row, one row above the shot number.
"""

from __future__ import annotations

import re
from typing import Any

from openpyxl.utils import get_column_letter

from .values import (
    is_range,
    number,
    parse_position,
    parse_time_window,
    parse_valve,
    parse_value,
)

#: Diagnostic labels whose value is a trigger time or window. Other labels in
#: the block (FastCam frame rate, filter, MD hardware notes, SEED) carry
#: settings, not times, and are kept as settings rather than guessed at.
DIAGNOSTIC_TRIGGERS = frozenset({
    "TS", "PIG", "IDS", "IDS(CES)", "CES", "HXR",
    "SXR", "SXR1", "SXR2", "SXR3", "SXR4", "IF", "IF H", "IF V",
})
PROBE_POSITIONS = frozenset({"TP", "TP(M/U)"})

#: NBI labels whose value is a time, and the one whose value is a duration.
NBI_START_LABELS = ("t0", "t1")
NBI_DURATION_LABELS = ("dt",)
#: HI labels whose value is a time (ms).
HI_TIME_PREFIXES = ("pfnspark", "tgas", "tinjection")
#: Cells that are a unit written beside a label, never a value.
UNIT_TOKENS = frozenset({"v", "ms", "bar", "s", "kv", "a", "ka", "kw", "-"})

GAS_SIDES = {"low field side": "LFS", "high field side": "HFS"}

_SECTION_PREFIXES = {
    "magnetic": ("magnetic coil",),
    "diagnostics": ("diagnostics trigger",),
    "heating": ("h & cd trigger",),
    "gas": ("gas injection",),
    "ip_peak": ("ip, maximum", "ip (ka)"),
    "ip_length": ("ip, pulse length", "ip (ms)"),
}
_UNIT_SUFFIX = re.compile(r"\s*\([^)]*\)\s*$")
SCAN_MAX_COLUMNS = 40


def _text(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value if value is not None else "")).strip()


def _key(value: Any) -> str:
    return _text(value).lower()


def _coordinate(row: int, column: int) -> str:
    return f"{get_column_letter(column)}{row}"


def _is_label(value: Any) -> bool:
    """A cell that names something rather than giving a value."""
    if value is None or isinstance(value, (int, float)):
        return False
    text = _text(value)
    return bool(text) and parse_time_window(text) is None and parse_valve(text) is None


def _label_stem(label: str) -> str:
    """``"PFNspark (ms)"`` -> ``"pfnspark"``: the name without its unit."""
    return _UNIT_SUFFIX.sub("", label).strip().lower()


def is_header_row(ws: Any, row: int) -> bool:
    """A card's first header row, recognised by its section titles.

    Not by ``Shot#`` in column A: operators overtype that cell (``C`` on
    2023-06-24), and a card whose header goes unrecognised loses every trigger.
    """
    if _key(ws.cell(row, 1).value) == "shot#":
        return True
    titles = {_key(ws.cell(row, column).value) for column in range(2, SCAN_MAX_COLUMNS + 1)}
    return any(title.startswith("diagnostics trigger") for title in titles) and any(
        title.startswith("h & cd trigger") for title in titles
    )


def find_header_row(ws: Any, shot_row: int) -> int | None:
    """The card's first header row within three rows above the shot."""
    for row in range(shot_row - 1, max(0, shot_row - 4), -1):
        if is_header_row(ws, row):
            return row
    return None


def section_columns(ws: Any, header_row: int) -> dict[str, int]:
    columns: dict[str, int] = {}
    for column in range(1, SCAN_MAX_COLUMNS + 1):
        text = _key(ws.cell(header_row, column).value)
        if not text:
            continue
        for name, prefixes in _SECTION_PREFIXES.items():
            if name not in columns and text.startswith(prefixes):
                columns[name] = column
    return columns


def _entry(system: str, identifier: str, raw: Any, row: int, column: int,
           window: tuple[Any, Any] | None, **detail: Any) -> dict[str, Any]:
    entry: dict[str, Any] = {
        "system": system,
        "identifier": identifier,
        "raw": _text(raw),
        "window_ms": list(window) if window is not None else None,
        "cell": _coordinate(row, column),
        "validation_flags": [] if window is not None else ["not_a_time_window"],
    }
    entry.update({key: value for key, value in detail.items() if value is not None})
    return entry


def _setting(section: str, label: str, raw: Any, row: int, column: int) -> dict[str, Any]:
    return {"section": section, "label": label, "raw": _text(raw), "cell": _coordinate(row, column)}


def _diagnostics(ws, rows, first, last, card) -> None:
    for label_column in range(first, last, 2):
        value_column = label_column + 1
        for row in rows:
            label_raw = ws.cell(row, label_column).value
            if not isinstance(label_raw, str) or not label_raw.strip():
                continue
            value = ws.cell(row, value_column).value
            if value is None or not _text(value):
                continue
            label = _text(label_raw).upper()
            if label in DIAGNOSTIC_TRIGGERS:
                card["timing"].append(_entry(
                    "diagnostic", label, value, row, value_column, parse_time_window(value)
                ))
            elif label in PROBE_POSITIONS:
                card["probe_positions"][label] = {
                    "raw": _text(value),
                    "positions_m": parse_position(value),
                    "cell": _coordinate(row, value_column),
                }
            else:
                card["settings"].append(_setting("diagnostics", _text(label_raw), value, row, value_column))


def _value_right_of(ws, row, first, last):
    """First non-unit cell right of a label, within its sub-block."""
    for column in range(first, last):
        value = ws.cell(row, column).value
        if value is None or not _text(value):
            continue
        if _key(value) in UNIT_TOKENS:
            continue
        return value, column
    return None, None


def _heating(ws, header2, rows, first, last, card) -> None:
    # Sub-blocks are named on the second header row: ECH, NBI, HI.
    starts: dict[str, int] = {}
    for column in range(first, last):
        text = _key(ws.cell(header2, column).value)
        if text.startswith("ech") and "ech" not in starts:
            starts["ech"] = column
        elif text == "nbi":
            starts["nbi"] = column
        elif text == "hi":
            starts["hi"] = column
    ech = starts.get("ech", first)
    nbi = starts.get("nbi")
    hi = starts.get("hi")

    # ECH: a window in the ECH column belongs to the nearest label above it,
    # the header "ECH (2.45 GHz)" included. Under an "HV"/"RF" label the cell
    # is a power-supply setting, and a lone number is not claimed as a time.
    label = _text(ws.cell(header2, ech).value) or None
    for row in rows:
        value = ws.cell(row, ech).value
        if value is None or not _text(value):
            continue
        if _is_label(value):
            label = _text(value)
            continue
        if label and label.lower().startswith("ech") and is_range(value):
            card["timing"].append(_entry("ec", label, value, row, ech, parse_time_window(value)))
        else:
            card["settings"].append(_setting("ec", label or "", value, row, ech))

    for name, start, stop in (
        ("nbi", nbi, hi if hi is not None else last),
        ("hi", hi, last),
    ):
        if start is None:
            continue
        for row in rows:
            label_raw = ws.cell(row, start).value
            if not _is_label(label_raw):
                continue
            label = _text(label_raw)
            value, column = _value_right_of(ws, row, start + 1, stop)
            if value is None:
                continue
            stem = _label_stem(label)
            if name == "nbi" and stem in NBI_START_LABELS:
                card["timing"].append(_entry("nbi", f"NBI {label}", value, row, column,
                                             parse_time_window(value)))
            elif name == "nbi" and stem in NBI_DURATION_LABELS:
                try:
                    card["nbi_duration_ms"] = number(_text(value))
                except ValueError:
                    card["settings"].append(_setting(name, label, value, row, column))
            elif name == "hi" and stem.startswith(HI_TIME_PREFIXES):
                card["timing"].append(_entry("hi", f"HI {label}", value, row, column,
                                             parse_time_window(value)))
            else:
                card["settings"].append(_setting(name, label, value, row, column))


def _gas(ws, rows, valve_column, card) -> None:
    trigger_column = valve_column + 1
    side: str | None = None
    for row in rows:
        valve_raw = ws.cell(row, valve_column).value
        trigger_raw = ws.cell(row, trigger_column).value
        valve = None
        if valve_raw is not None and _text(valve_raw):
            if _is_label(valve_raw):
                text = _text(valve_raw)
                side = GAS_SIDES.get(text.lower(), text)
            else:
                valve = parse_valve(valve_raw)
        identifier = side or "gas"
        if trigger_raw is not None and _text(trigger_raw):
            card["timing"].append(_entry(
                "gas", identifier, trigger_raw, row, trigger_column,
                parse_time_window(trigger_raw),
                valve_voltage_V=(valve or {}).get("voltage_V"),
                species=(valve or {}).get("species"),
                valve_cell=_coordinate(row, valve_column) if valve else None,
            ))
        elif valve is None and is_range(valve_raw):
            # Older cards log the window in the valve column itself.
            card["timing"].append(_entry(
                "gas", identifier, valve_raw, row, valve_column, parse_time_window(valve_raw)
            ))
        elif valve is not None:
            card["settings"].append(_setting("gas", identifier, valve_raw, row, valve_column))


def extract_card(ws: Any, shot_row: int, end_row: int) -> dict[str, Any] | None:
    """Return the structured content of the card whose shot number is on ``shot_row``.

    ``None`` when the rows above are not a card header, so the caller can fall
    back to the generic extraction instead of misreading an irregular block.
    Times stay on the ShotLog clock (``window_ms``); no offset is applied here.
    """
    header1 = find_header_row(ws, shot_row)
    if header1 is None:
        return None
    header2 = header1 + 1
    columns = section_columns(ws, header1)
    if "diagnostics" not in columns or "heating" not in columns:
        return None
    # A shot block runs to the next shot number, which puts the *next* card's
    # two header rows -- and the next shot's plasma current -- inside it. It
    # also swallows any card whose column A is not a shot number, such as a
    # reference card ("REF!!! 43013"), whose settings are not this shot's.
    last = end_row
    for row in range(shot_row + 1, end_row + 1):
        if is_header_row(ws, row):
            last = row - 1
            break
    rows = range(shot_row, last + 1)
    card: dict[str, Any] = {
        "header_rows": [header1, header2],
        "rows": [shot_row, last],
        "timing": [],
        "probe_positions": {},
        "settings": [],
        "plasma_current": {},
        "remarks": [],
    }
    gas_column = columns.get("gas")
    heating_end = gas_column or columns.get("ip_peak") or SCAN_MAX_COLUMNS
    _diagnostics(ws, rows, columns["diagnostics"], columns["heating"], card)
    _heating(ws, header2, rows, columns["heating"], heating_end, card)
    if gas_column is not None:
        _gas(ws, rows, gas_column, card)

    for name, unit in (("ip_peak", "kA"), ("ip_length", "ms")):
        column = columns.get(name)
        if column is None:
            continue
        value = ws.cell(header2, column).value
        if value is None or not _text(value):
            continue
        record = parse_value(value, unit)
        record["source"] = {"cells": [_coordinate(header2, column)]}
        card["plasma_current"]["peak" if name == "ip_peak" else "pulse_length"] = record

    remark_column = columns.get("ip_peak")
    if remark_column is not None:
        for row in rows:
            value = ws.cell(row, remark_column).value
            if value is not None and _text(value) and _key(value) not in {"remark", "remarks"}:
                card["remarks"].append(_text(value))
    return card
