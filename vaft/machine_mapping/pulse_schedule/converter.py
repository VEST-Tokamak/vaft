"""Convert one ShotLog worksheet into a session document.

A session is one worksheet: one operating day, split into run groups (the
operator's titled experiment blocks) and shots. Every structured value keeps
the cell it came from, and every shot keeps all of its raw cells in
``source_occurrences``, so nothing the parser declined to interpret is lost.

Ported from the standalone ``VEST_ShotLog`` converter (#995). The output
omits the conversion timestamp on purpose: a session document is a pure
function of the workbook bytes and the schema registry, so reconverting an
unchanged workbook reproduces it byte for byte and FileDB products do not
churn.
"""

from __future__ import annotations

from copy import deepcopy
from dataclasses import dataclass
import hashlib
import json
from pathlib import Path
import re
from typing import Any
import unicodedata

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
import yaml

from .card import extract_card
from .values import normalize_status, parse_value

SHOT_RE = re.compile(r"^\s*(\d{3,6})\s*$")
SHOT_RANGE_RE = re.compile(r"^\s*(\d{3,6})\s*[-~–]\s*(\d{3,6})\s*$")
#: ``Ref. #41234``, ``Reference: 41234``, and the ``REF!!! 43013`` a reference
#: card actually carries -- up to six non-digit characters between the word and
#: the number.
REFERENCE_RE = re.compile(r"\bref(?:erence)?(?![a-z])[^\d\n]{0,6}?\s*(\d{3,6})\b", re.IGNORECASE)
DATE_RE = re.compile(r"^(\d{2})(\d{2})(\d{2})")
DATE8_RE = re.compile(r"^(\d{4})(\d{2})(\d{2})")
DATE_ISO_RE = re.compile(r"^(\d{4})-(\d{2})-(\d{2})")  # ERC_ShotLog names sheets 2013-01-21
FILE_DATE_RE = re.compile(r"ShotLog_(\d{4})_(\d{1,2})", re.IGNORECASE)
SCAN_MAX_COLUMNS = 128
#: A shot range longer than this is a typo, not a run of identical shots.
MAX_SHOT_RANGE = 200
#: Labels that head a column but are never its value.
STRUCTURAL_LABELS = frozenset({"remark", "remarks", "fail", "shot#", "shot"})


@dataclass(frozen=True)
class Marker:
    row: int
    kind: str
    shots: tuple[int, ...] = ()
    reference: int | None = None
    references: tuple[int, ...] = ()


def _normalise(text: Any) -> str:
    return re.sub(r"\s+", " ", str(text if text is not None else "").replace("\n", " ")).strip().lower()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def derive_session(sheet_name: str, source: Path) -> tuple[str | None, str]:
    """``(experiment_date, session_id)`` from a sheet named ``YYMMDD`` or ``YYYYMMDD``."""
    match8 = DATE8_RE.match(sheet_name) or DATE_ISO_RE.match(sheet_name)
    match6 = DATE_RE.match(sheet_name)
    if match8:
        year, month, day = match8.groups()
        date = f"{year}-{month}-{day}"
    elif match6:
        year, month, day = match6.groups()
        date = f"20{year}-{month}-{day}"
    else:
        file_match = FILE_DATE_RE.search(source.name)
        date = f"{file_match.group(1)}-{int(file_match.group(2)):02d}-01" if file_match else None
    suffix = re.sub(r"[^A-Za-z0-9가-힣_-]+", "-", sheet_name).strip("-") or "sheet"
    return date, f"{date or 'undated'}__{suffix}"


def _marker(value: Any, row: int) -> Marker | None:
    raw = str(value if value is not None else "")
    text = raw.replace("\n", " ").strip()
    if REFERENCE_RE.search(text):
        # "ref.\n37911 38740" names two; every number after the word counts.
        tail = text[REFERENCE_RE.search(text).start():]
        numbers = tuple(int(number) for number in re.findall(r"\d{3,6}", tail))
        return Marker(row=row, kind="reference", reference=numbers[0], references=numbers)
    # The whole cell first: a range is often broken across lines ("4169\n~4173").
    parts = [text] if (SHOT_RE.match(text) or SHOT_RANGE_RE.match(text)) else re.split(r"[\n,]+", raw)
    shots: list[int] = []
    # Otherwise one cell can log several shots, one per line or comma-separated
    # ("39853\n39854\n39855"); each part must be a shot or a forward range.
    for part in parts:
        part = part.strip()
        if not part:
            continue
        exact = SHOT_RE.match(part)
        interval = SHOT_RANGE_RE.match(part)
        if exact:
            shots.append(int(exact.group(1)))
        elif interval:
            first, last = int(interval.group(1)), int(interval.group(2))
            if not (first <= last and last - first <= MAX_SHOT_RANGE):
                return None
            shots.extend(range(first, last + 1))
        else:
            return None
    if not shots:
        return None
    return Marker(row=row, kind="shots", shots=tuple(dict.fromkeys(shots)))


def marker_shots(value: Any) -> tuple[int, ...]:
    """The shots a column-A cell logs, or ``()`` when it logs none."""
    marker = _marker(value, 0)
    return marker.shots if marker is not None and marker.kind == "shots" else ()


def _within(shot: int, span: tuple[int, int | None] | None) -> bool:
    if span is None:
        return True
    lower, upper = span
    return shot >= lower and (upper is None or shot <= upper)


def _cell_record(cell: Any) -> dict[str, Any]:
    raw = cell.value
    return {
        "cell": f"{get_column_letter(cell.column)}{cell.row}",
        "raw": str(raw),
        "formula": str(raw) if isinstance(raw, str) and raw.startswith("=") else None,
    }


def _block_cells(ws: Any, start: int, end: int) -> list[dict[str, Any]]:
    return [
        _cell_record(cell)
        for row in ws.iter_rows(min_row=start, max_row=end, max_col=SCAN_MAX_COLUMNS)
        for cell in row
        if cell.value is not None and str(cell.value).strip()
    ]


def _title_between(ws: Any, start: int, end: int, column_a_only: bool = False) -> str | None:
    """Find a likely operator title without mistaking a repeated header for one.

    On the modern card the operator's title is always in column A; anything
    else between two cards (a reference card's ``VB``/``filter`` row) is not.
    """
    header_terms = ("shot", "tf", "p/s", "charging", "remark", "diagnostic", "gas injection")
    titles: list[str] = []
    for row in ws.iter_rows(min_row=max(1, start), max_row=max(0, end), max_col=SCAN_MAX_COLUMNS):
        values = [str(cell.value).strip() for cell in row if cell.value is not None and str(cell.value).strip()]
        if column_a_only:
            first = row[0].value if row else None
            if len(values) != 1 or first is None or str(first).strip() != values[0]:
                continue
        if len(values) > 2:
            continue
        for value in values:
            normalised = _normalise(value)
            if any(term in normalised for term in header_terms) or SHOT_RE.match(value):
                continue
            if len(value) > 5:
                titles.append(value)
    return titles[-1] if titles else None


def detect_schema(ws: Any, registry: dict[str, Any]) -> tuple[str | None, dict[str, Any]]:
    """Pick the era whose required headers all appear on the sheet.

    Ties go to the higher ``detection.priority``: the modern card also
    contains ``C (mF)``, ``NBI``, ``Gas Injection`` and ``Remark``, so the
    integrated_v3 test passes on it too. Counting headers does not separate
    them -- integrated_v3 asks for more, and more generic, ones.
    """
    values = [
        _normalise(cell.value)
        for row in ws.iter_rows(max_col=SCAN_MAX_COLUMNS)
        for cell in row
        if cell.value is not None
    ]
    text = "\n".join(values)
    candidates: list[tuple[float, int, str, dict[str, Any]]] = []
    for schema_id, schema in registry["schemas"].items():
        required = schema["detection"].get("required_headers", [])
        matched = sum(1 for header in required if _normalise(header) in text)
        score = matched / len(required) if required else 0.0
        priority = schema["detection"].get("priority", 0)
        candidates.append((score, priority, schema_id, schema))
    if not candidates:
        return None, {"score": 0.0}
    score, _, schema_id, schema = max(candidates, key=lambda item: item[:3])
    minimum = schema.get("detection", {}).get("minimum_score", 1.0)
    return (schema_id, schema) if score >= minimum else (None, {"score": score})


def _find_value_below(ws: Any, row: int, column: int, end: int, labels: set[str]) -> tuple[Any, str] | None:
    """The first value within two rows under a header, skipping labels.

    ``labels`` is every alias the schema knows plus the structural column
    labels. A cell equal to one of them heads something; it is never a value
    -- reading it as one is how ``"Remarks"`` became a plasma current.
    """
    for candidate_row in range(row + 1, min(end, row + 2) + 1):
        value = ws.cell(candidate_row, column).value
        if value is None or not str(value).strip():
            continue
        if _normalise(value) in labels:
            return None
        return value, f"{get_column_letter(column)}{candidate_row}"
    return None


def set_path(target: dict[str, Any], path: str, value: Any) -> None:
    parts = [part for part in path.replace("shots[].", "").split(".") if part]
    cursor = target
    for part in parts[:-1]:
        cursor = cursor.setdefault(part, {})
    cursor[parts[-1]] = value


def _schema_labels(schema: dict[str, Any]) -> set[str]:
    labels = {_normalise(alias) for field in schema["fields"] for alias in field["aliases"]}
    return labels | STRUCTURAL_LABELS


def _extract_mapped_fields(ws: Any, start: int, end: int, schema: dict[str, Any]) -> dict[str, Any]:
    labels = _schema_labels(schema)
    result: dict[str, Any] = {}
    for field in schema["fields"]:
        path = field["path"]
        if not path.startswith("shots[]."):
            continue
        aliases = {_normalise(alias) for alias in field["aliases"]}
        found: tuple[Any, str] | None = None
        for row in ws.iter_rows(min_row=start, max_row=end, max_col=SCAN_MAX_COLUMNS):
            for cell in row:
                if cell.value is not None and _normalise(cell.value) in aliases:
                    found = _find_value_below(ws, cell.row, cell.column, end, labels)
                    if found:
                        break
            if found:
                break
        if found:
            raw, source_cell = found
            parsed = parse_value(raw, field.get("unit"), field.get("parsing") == "trigger_range")
            parsed["source"] = {"cells": [source_cell]}
            set_path(result, path, parsed)
    return result


def _reference_entries(shots: list[int]) -> list[dict[str, Any]]:
    return [{"shot": shot, "relation": "reference", "source": {"cells": []}} for shot in shots]


def _remarks(block_cells: list[dict[str, Any]]) -> list[str]:
    terms = ("fail", "error", "remark", "remarks", "안나감", "나감", "이상", "not saved")
    headers = {"remark", "remarks", "fail"}
    return [
        cell["raw"]
        for cell in block_cells
        if _normalise(cell["raw"]) not in headers and any(term in cell["raw"].lower() for term in terms)
    ]


def _deep_merge(base: dict[str, Any], incoming: dict[str, Any]) -> dict[str, Any]:
    result = deepcopy(base)
    for key, value in incoming.items():
        if isinstance(value, dict) and isinstance(result.get(key), dict):
            result[key] = _deep_merge(result[key], value)
        else:
            result[key] = value
    return result


def _apply_card(record: dict[str, Any], card: dict[str, Any]) -> None:
    """Fold one card into a shot record; a later card for the same shot wins.

    The converter keeps repeated blocks for one shot in source-row order, and a
    later block is a logged correction or retry, so it supersedes per timing
    identifier rather than being appended as a second schedule.
    """
    planned = record["planned_configuration"]
    merged = {(item["system"], item["identifier"]): item for item in planned.get("timing", [])}
    for item in card["timing"]:
        merged[(item["system"], item["identifier"])] = item
    planned["timing"] = list(merged.values())
    planned.setdefault("probe_positions", {}).update(card["probe_positions"])
    planned["settings"] = planned.get("settings", []) + card["settings"]
    if "nbi_duration_ms" in card:
        planned["heating_and_current_drive"]["nbi_duration_ms"] = card["nbi_duration_ms"]
    outcome = record["observed_outcome"]
    outcome["plasma_current"].update(card["plasma_current"])
    for remark in card["remarks"]:
        if remark not in outcome["remarks"]:
            outcome["remarks"].append(remark)


def _new_shot(shot: int, cells: list[dict[str, Any]], mapped: dict[str, Any]) -> dict[str, Any]:
    record: dict[str, Any] = {
        "shot": shot,
        "status": normalize_status([cell["raw"] for cell in cells]),
        "planned_configuration": {
            "magnetic": {"tf": None, "power_supplies": []},
            "heating_and_current_drive": {},
            "gas_injection": {},
            "diagnostics": {},
        },
        "observed_outcome": {"plasma_current": {}, "remarks": _remarks(cells)},
        "source_occurrences": [{"cells": cells}],
    }
    for root, value in mapped.items():
        if isinstance(value, dict) and isinstance(record.get(root), dict):
            record[root] = _deep_merge(record[root], value)
        else:
            record[root] = value
    return record


def load_overrides(path: Path | None) -> tuple[list[dict[str, Any]], str | None]:
    """Read a review-overrides document (YAML; JSON is a subset)."""
    if not path or not path.exists():
        return [], None
    try:
        overrides = yaml.safe_load(path.read_text(encoding="utf-8")) or []
    except yaml.YAMLError as error:
        return [], f"invalid_override_document:{error}"
    if not isinstance(overrides, list):
        return [], "invalid_override_document:not_a_list"
    return overrides, None


def _apply_overrides(dataset: dict[str, Any], overrides_path: Path | None) -> None:
    review: dict[str, Any] = {"applied_overrides": [], "override_conflicts": []}
    dataset["review"] = review
    overrides, error = load_overrides(overrides_path)
    if error:
        review["override_conflicts"].append({"reason": error})
        return
    for override in overrides:
        target = override.get("target", {})
        if override.get("source_sha256") != dataset["provenance"]["source_sha256"]:
            continue
        if target.get("session_id") != dataset["experiment_session"]["id"]:
            continue
        selected = None
        for group in dataset["run_groups"]:
            if group["id"] != target.get("run_group_id"):
                continue
            selected = next((shot for shot in group["shots"] if shot["shot"] == target.get("shot")), None)
            break
        if selected is None:
            review["override_conflicts"].append({"id": override.get("id"), "reason": "target_not_found"})
            continue
        path = override.get("path", "")
        if not path:
            review["override_conflicts"].append({"id": override.get("id"), "reason": "missing_path"})
            continue
        set_path(selected, path, override.get("value"))
        review["applied_overrides"].append(override.get("id"))


def validate_dataset(dataset: dict[str, Any], registry: dict[str, Any]) -> list[str]:
    errors: list[str] = []
    if dataset.get("schema_version") not in {*registry["schemas"], "unclassified"}:
        errors.append("unknown_schema_version")
    for key in ("provenance", "experiment_session", "run_groups"):
        if key not in dataset:
            errors.append(f"missing:{key}")
    if '"shot_no"' in json.dumps(dataset, ensure_ascii=False, default=str):
        errors.append("prohibited_key:shot_no")
    for group in dataset.get("run_groups", []):
        seen: set[int] = set()
        for shot in group.get("shots", []):
            value = shot.get("shot")
            if not isinstance(value, int):
                errors.append(f"invalid_shot:{value}")
            elif value in seen:
                errors.append(f"duplicate_shot:{value}")
            seen.add(value)
    return errors


def convert_sheet(
    source: Path,
    sheet_name: str,
    registry: dict[str, Any],
    overrides_path: Path | None = None,
    workbook: Any | None = None,
    source_hash: str | None = None,
    span: tuple[int, int | None] | None = None,
) -> dict[str, Any]:
    """Convert ``sheet_name`` of ``source`` into a session document.

    ``span`` is the range of shots the workbook can hold (its ``#first-last``
    with a margin). A column-A number outside it is not one of this sheet's
    shots -- a reference to an older discharge, or a value that only looks
    like a shot number -- and is recorded as such instead of becoming a record.
    """
    source = Path(source)
    owns_workbook = workbook is None
    if workbook is None:
        workbook = load_workbook(source, read_only=False, data_only=False)
    try:
        return _convert(source, sheet_name, registry, overrides_path, workbook, source_hash, span)
    finally:
        if owns_workbook:
            workbook.close()


def _convert(source, sheet_name, registry, overrides_path, workbook, source_hash, span=None) -> dict[str, Any]:
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Sheet not found: {sheet_name}")
    ws = workbook[sheet_name]
    provenance = {
        "source_file": unicodedata.normalize("NFC", source.name),
        "source_sha256": source_hash or sha256_file(source),
        "sheet_name": sheet_name,
    }
    schema_id, schema = detect_schema(ws, registry)
    classification = None
    if schema_id is None:
        # No template matched, but the sheet may still log shots. They keep
        # their raw cells under an empty field list rather than vanishing.
        classification = schema
        schema = {"fields": [], "version": None}

    experiment_date, session_id = derive_session(sheet_name, source)
    card_extractor = schema.get("extractor") == "modern_card"
    markers = [marker for row in range(1, ws.max_row + 1) if (marker := _marker(ws.cell(row, 1).value, row))]
    marked = {marker.row for marker in markers}
    out_of_span: list[dict[str, Any]] = []
    if span is not None:
        kept: list[Marker] = []
        for marker in markers:
            if marker.kind != "shots":
                kept.append(marker)
                continue
            inside = tuple(shot for shot in marker.shots if _within(shot, span))
            outside = [shot for shot in marker.shots if not _within(shot, span)]
            if outside:
                out_of_span.append({"cell": f"A{marker.row}", "shots": outside})
            if inside:
                kept.append(Marker(row=marker.row, kind="shots", shots=inside))
            else:
                # Its card, if any, is an older shot's settings: keep it as a
                # reference, the way an explicit "Ref. #N" card is kept.
                kept.append(Marker(row=marker.row, kind="reference", reference=outside[0],
                                   references=tuple(outside)))
        markers = kept
    # A column-A cell of digits and separators that is not a marker is a shot
    # number the parser could not read (a reversed range, "43782/3"). Its card
    # is lost, so it is reported for review rather than skipped in silence.
    unreadable = [
        {"cell": f"A{row}", "raw": str(value)}
        for row in range(1, ws.max_row + 1)
        if row not in marked
        and (value := ws.cell(row, 1).value) is not None
        and re.fullmatch(r"[\d\s,~/\-\u2013]+", str(value))
        and re.search(r"\d{3,6}", str(value))
    ]
    groups: list[dict[str, Any]] = []
    active_group: dict[str, Any] | None = None
    previous_end = 0
    references: list[int] = []

    for marker in markers:
        if marker.kind == "reference" and marker.reference is not None:
            references.extend(marker.references or (marker.reference,))
            continue
        next_rows = [candidate.row for candidate in markers if candidate.row > marker.row]
        end = (min(next_rows) - 1) if next_rows else ws.max_row
        title = _title_between(ws, previous_end + 1, marker.row - 1, column_a_only=card_extractor)
        if active_group is None or title:
            group_index = len(groups) + 1
            active_group = {
                "id": f"{session_id}__rg-{group_index:03d}",
                "project_index": group_index,
                "title": {"raw": title},
                "segmentation": {"method": "automatic", "confidence": "medium", "basis": "title_gap_or_first_shot"},
                "references": [],
                "shots": [],
            }
            groups.append(active_group)
        # A reference belongs to the run group of the next shot. Flushing only
        # when a new group opened dropped every reference logged mid-group.
        active_group["references"].extend(_reference_entries(references))
        references = []
        card = extract_card(ws, marker.row, end) if card_extractor else None
        block_end = card["rows"][1] if card else end
        cells = _block_cells(ws, marker.row, block_end)
        # Modern templates place their two-level header immediately before a shot row.
        mapped = _extract_mapped_fields(ws, max(1, marker.row - 2), block_end, schema)
        existing = {shot["shot"]: shot for shot in active_group["shots"]}
        for shot_number in marker.shots:
            if shot_number in existing:
                item = existing[shot_number]
                item["source_occurrences"].append({"cells": cells})
            else:
                item = _new_shot(shot_number, cells, mapped)
                active_group["shots"].append(item)
                existing[shot_number] = item
            if card is not None:
                _apply_card(item, card)
        # Where the next title search starts. `end` runs to the next shot
        # number, which left every later search an empty range and put a whole
        # sheet in one run group. A card knows where it ends; other templates
        # keep the old span (their title rows are not yet told apart from
        # remark rows -- #995 increment 2).
        previous_end = block_end

    if groups and references:
        groups[-1]["references"].extend(_reference_entries(references))
        references = []
    if not groups:
        groups = [{
            "id": f"{session_id}__rg-001",
            "project_index": 1,
            "title": {"raw": _title_between(ws, 1, ws.max_row)},
            "segmentation": {"method": "automatic", "confidence": "low", "basis": "no_physical_shot_marker"},
            "references": _reference_entries(references),
            "shots": [],
        }]

    dataset = {
        "schema_version": schema_id or "unclassified",
        "schema_revision": schema["version"],
        "provenance": provenance,
        "experiment_session": {"id": session_id, "experiment_date": experiment_date, "session_label": sheet_name},
        "run_groups": groups,
    }
    _apply_overrides(dataset, overrides_path)
    if unreadable:
        dataset["review"]["unreadable_shot_cells"] = unreadable
    if out_of_span:
        dataset["review"]["out_of_span_markers"] = out_of_span
    if classification is not None:
        dataset["classification"] = classification
        dataset["review"]["reason"] = "unclassified_template"
    errors = validate_dataset(dataset, registry)
    if errors:
        dataset["review"]["validation_errors"] = errors
    return dataset


def session_path(output_dir: Path, dataset: dict[str, Any]) -> Path:
    session = dataset["experiment_session"]["id"]
    year = (dataset["experiment_session"].get("experiment_date") or "undated")[:4]
    return Path(output_dir) / "experiment-days" / year / f"{session}.yaml"
